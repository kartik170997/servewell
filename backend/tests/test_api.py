import io
import json
import os
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pytest
import httpx
from fastapi.testclient import TestClient
import main

@pytest.mark.parametrize('url,provider,key',[
    ('https://api.anthropic.com/v1/messages','Claude','CLAUDE_API_KEY'),
    ('https://api.pinecone.io/embed','Pinecone','PINECONE_API_KEY'),
    ('https://test.svc.pinecone.io/query','Pinecone','PINECONE_API_KEY'),
])
def test_auth_failure_identifies_provider_without_exposing_key(monkeypatch,url,provider,key):
    monkeypatch.setattr(httpx.Client,'request',lambda *a,**kw:httpx.Response(401,json={'error':'SECRET_PROVIDER_RESPONSE'}))
    with pytest.raises(ValueError) as error:
        main.request('POST',url,{'Api-Key':'SECRET_KEY'})
    message=str(error.value)
    assert provider in message and key in message and '401' in message
    assert 'SECRET' not in message

@pytest.fixture
def client(tmp_path,monkeypatch):
    for names in main.SETTINGS_ENV.values():
        for name in names: monkeypatch.delenv(name, raising=False)
    monkeypatch.setattr(main,'DATA',tmp_path)
    monkeypatch.setenv('APP_TOKEN','test-workspace-token-long-enough')
    monkeypatch.setenv('DISABLE_WORKER','1')
    with TestClient(main.app) as c:
        c.headers['Authorization']='Bearer test-workspace-token-long-enough'
        yield c

@pytest.fixture
def providers(monkeypatch):
    calls=[]
    def fake_request(method,url,headers,**kw):
        body=kw.get('json',{}); calls.append((url,body))
        if url.endswith('/v1/messages'):
            return {'content':[{'type':'tool_use','name':'record_document','input':{'title':'RAG handbook','summary':'Hybrid retrieval combines BM25 and vectors.','tags':['engineering','search'],'language':'English','document_type':'Guide'}}]}
        if url.endswith('/embed'): return {'data':[{'values':[0.1,0.2,0.3]} for t in body['inputs']]}
        if url.endswith('/vectors/upsert'): return {'upsertedCount':len(body['vectors'])}
        if url.endswith('/vectors/delete'): return {}
        if url.endswith('/query'):
            with main.db() as db:
                rows=db.execute('SELECT id FROM chunks ORDER BY position DESC').fetchall()
            return {'matches':[{'id':r['id'],'score':0.8} for r in rows]}
        if url.endswith('/rerank'):
            return {'data':[{'index':i,'score':.99-i*.1} for i in list(range(len(body['documents'])))[::-1][:body['top_n']]]}
        if url.endswith('/describe_index_stats'): return {'dimension':3}
        if '/v1/models/' in url:return {'id':'test-model'}
        raise AssertionError('Unexpected upstream request: '+url)
    monkeypatch.setattr(main,'request',fake_request)
    return calls

def configure(client):
    s=main.Settings(claude_key='SECRET_CLAUDE',pinecone_key='SECRET_PINECONE',index_host='https://test.svc.pinecone.io')
    r=client.put('/api/settings',json=s.model_dump()); assert r.status_code==200
    return s

def add_doc(client,content=b'Hybrid retrieval combines BM25 keyword search with dense vectors. '*60,name='guide.md'):
    r=client.post('/api/documents',files=[('files',(name,content,'text/plain'))]); assert r.status_code==202
    return r.json()['documents'][0]['id']

def test_auth_and_secret_storage(client):
    assert client.get('/api/settings',headers={'Authorization':'Bearer bad'}).status_code==401
    configure(client)
    r=client.get('/api/settings')
    assert r.json()['claude_configured'] is True
    assert 'SECRET' not in r.text
    assert b'SECRET' not in (main.DATA/'folio.db').read_bytes()
    assert b'SECRET' not in (main.DATA/'folio.db-wal').read_bytes() if (main.DATA/'folio.db-wal').exists() else True
    s=r.json();s.pop('claude_configured');s.pop('pinecone_configured')
    assert client.put('/api/settings',json=s).status_code==200
    assert main.settings().claude_key=='SECRET_CLAUDE'

def test_environment_settings_without_saved_data(client,monkeypatch,tmp_path):
    from dotenv import load_dotenv
    env=tmp_path/'.env'
    env.write_text('CLAUDE_API_KEY=ENV_CLAUDE\npinecone_key=ENV_PINECONE\nPINECONE_INDEX_HOST=example.svc.pinecone.io\nPINECONE_NAMESPACE=env-space\nEMBEDDING_MODEL=env-embedding\n')
    # Register these names with monkeypatch so loading the file cannot leak across tests.
    for name in ('CLAUDE_API_KEY','pinecone_key','PINECONE_INDEX_HOST','PINECONE_NAMESPACE','EMBEDDING_MODEL'):
        monkeypatch.setenv(name,'')
        monkeypatch.delenv(name)
    load_dotenv(env,override=False)
    s=main.settings()
    assert s.claude_key=='ENV_CLAUDE' and s.pinecone_key=='ENV_PINECONE'
    assert s.index_host=='https://example.svc.pinecone.io'
    assert s.namespace=='env-space' and s.embedding_model=='env-embedding'
    r=client.get('/api/settings')
    assert r.json()['claude_configured'] and r.json()['pinecone_configured']
    assert 'ENV_CLAUDE' not in r.text and 'ENV_PINECONE' not in r.text
    assert client.post('/api/documents',files=[('files',('env.md',b'Environment configured upload','text/plain'))]).status_code==202

def test_saved_values_override_environment_and_blank_values_fall_back(client,monkeypatch):
    monkeypatch.setenv('CLAUDE_API_KEY','ENV_CLAUDE')
    monkeypatch.setenv('PINECONE_API_KEY','ENV_PINECONE')
    monkeypatch.setenv('PINECONE_INDEX_HOST','env.svc.pinecone.io')
    with main.db() as c:
        raw={'claude_key':'SAVED_CLAUDE','pinecone_key':'','namespace':'saved-space'}
        c.execute('INSERT INTO settings VALUES (1,?)',(main.cipher().encrypt(json.dumps(raw).encode()),))
    s=main.settings()
    assert s.claude_key=='SAVED_CLAUDE' and s.pinecone_key=='ENV_PINECONE'
    assert s.namespace=='saved-space' and s.index_host=='https://env.svc.pinecone.io'
    assert client.put('/api/settings',json={'claude_key':'','pinecone_key':'','rerank_model':'custom-rerank'}).status_code==200
    assert main.settings().claude_key=='SAVED_CLAUDE'
    assert main.saved_settings()['pinecone_key']==''
    monkeypatch.setenv('PINECONE_API_KEY','ROTATED_PINECONE')
    assert main.settings().pinecone_key=='ROTATED_PINECONE'
    assert client.put('/api/settings',json={'pinecone_key':'SAVED_PINECONE'}).status_code==200
    assert main.settings().pinecone_key=='SAVED_PINECONE'

def test_environment_alias_priority_and_builtin_defaults(client,monkeypatch):
    assert main.settings()==main.Settings()
    monkeypatch.setenv('ANTHROPIC_API_KEY','ANTHROPIC_VALUE')
    monkeypatch.setenv('pinecone_key','LOWERCASE_VALUE')
    assert main.settings().claude_key=='ANTHROPIC_VALUE'
    assert main.settings().pinecone_key=='LOWERCASE_VALUE'
    monkeypatch.setenv('PINECONE_API_KEY','CANONICAL_VALUE')
    assert main.settings().pinecone_key=='CANONICAL_VALUE'

def test_ingestion_duplicate_search_filters_delete(client,providers):
    configure(client); doc=add_doc(client); main.ingest(doc)
    d=client.get('/api/documents/'+doc).json()
    assert d['status']=='ready' and d['chunk_count']==1
    assert len(d['chunks'])==d['chunk_count']
    assert d['tags']==['engineering','search']
    assert add_doc(client)==doc
    for mode in ('lexical','vector','hybrid'):
        r=client.post('/api/search',json={'query':'hybrid retrieval','mode':mode,'rerank':False}).json()
        assert r['results'],mode
        assert all(x['document_id']==doc for x in r['results'])
    assert not client.post('/api/search',json={'query':'hybrid','mode':'lexical','rerank':False,'tag':'missing'}).json()['results']
    assert client.delete('/api/documents/'+doc).status_code==200
    assert any(url.endswith('/vectors/delete') for url,body in providers)
    assert not client.get('/api/documents').json()
    assert client.get('/api/documents/'+doc).status_code==404
    with main.db() as db: assert db.execute('SELECT count(*) FROM chunks_fts').fetchone()[0]==0

def test_failed_docs_invisible_retry_idempotent(client,providers,monkeypatch):
    configure(client);doc=add_doc(client)
    original=main.embed
    monkeypatch.setattr(main,'embed',lambda *args: (_ for _ in ()).throw(ValueError('Quota exceeded')))
    main.ingest(doc)
    assert main.get_doc(doc)['status']=='failed'
    assert not client.post('/api/search',json={'query':'hybrid','mode':'lexical','rerank':False}).json()['results']
    monkeypatch.setattr(main,'embed',original)
    assert client.post('/api/documents/'+doc+'/retry').status_code==200
    main.ingest(doc)
    d=main.get_doc(doc);assert d['status']=='ready'
    with main.db() as db: assert db.execute('SELECT count(*) FROM chunks').fetchone()[0]==d['chunk_count']
    assert client.post('/api/documents/'+doc+'/retry').status_code==409

def test_settings_guards_and_host_validation(client):
    s=configure(client);doc=add_doc(client)
    s.embedding_model='new-model'
    assert client.put('/api/settings',json=s.model_dump()).status_code==409
    assert client.delete('/api/documents/'+doc).status_code==409
    for host in ('http://127.0.0.1','https://evil.com','https://test.pinecone.io@evil.com','https://test.pinecone.io/path','https://test.pinecone.io:444'):
        s.index_host=host
        assert client.put('/api/settings',json=s.model_dump()).status_code==422

def test_rerank_mapping_and_rrf(client,providers):
    configure(client);doc=add_doc(client);main.ingest(doc)
    other=add_doc(client,content=b'A second document on hybrid retrieval.',name='second.md');main.ingest(other)
    before=client.post('/api/search',json={'query':'hybrid','mode':'lexical','rerank':False,'top_k':20}).json()['results']
    after=client.post('/api/search',json={'query':'hybrid','mode':'lexical','rerank':True,'top_k':20}).json()['results']
    assert [x['id'] for x in after]==[x['id'] for x in before][::-1]
    assert after[0]['score_type']=='Rerank'
    rows=main.rrf([('lexical',[{'id':'a'},{'id':'b'}]),('vector',[{'id':'b'},{'id':'c'}])])
    assert rows[0]['id']=='b' and rows[0]['sources']==['lexical','vector']

def test_parse_docx_html_chunk_coverage(tmp_path):
    from docx import Document
    p=tmp_path/'sample.docx';d=Document();d.add_paragraph('Before');d.add_table(rows=1,cols=1).cell(0,0).text='Middle';d.add_paragraph('After');d.save(p)
    pages,text=main.parse_file(p,'.docx');assert text=='Before\nMiddle\nAfter'
    p.write_text('<script>evil()</script><h1>Good</h1><p>Content</p>')
    _,text=main.parse_file(p,'.html');assert 'evil' not in text and 'Content' in text
    original=' '.join('word'+str(i) for i in range(400))
    chunks=main.chunk_pages([{'page':2,'text':original}],200,30)
    assert all(c['page']==2 and len(c['text'])<=200 for c in chunks)
    assert all('word'+str(i) in ' '.join(c['text'] for c in chunks) for i in range(400))

def test_upload_errors_and_query_validation(client):
    configure(client)
    r=client.post('/api/documents',files=[('files',('bad.exe',b'bad','application/octet-stream'))])
    assert r.status_code==202 and r.json()['documents'][0]['error']
    assert client.post('/api/search',json={'query':' ','mode':'lexical'}).status_code==422
    assert client.post('/api/search',json={'query':'x','top_k':20,'candidates':5}).status_code==422

def test_connections(client,providers):
    configure(client)
    r=client.post('/api/settings/test').json()
    assert all(x['ok'] for x in r.values())


def test_validation_does_not_echo_secrets(client):
    r=client.put('/api/settings',json={'claude_key':'SECRET'*100})
    assert r.status_code==422 and 'SECRET' not in r.text


def test_pdf_text_and_page_citations(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    writer=PdfWriter()
    for sentence in ('Hybrid search handbook','Pinecone vector retrieval'):
        page=writer.add_blank_page(width=612,height=792)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        stream=DecodedStreamObject();stream.set_data(f'BT /F1 12 Tf 50 700 Td ({sentence}) Tj ET'.encode())
        page[NameObject('/Contents')]=writer._add_object(stream)
    p=tmp_path/'handbook.pdf';writer.write(p)
    pages,text=main.parse_file(p,'.pdf')
    assert len(pages)==2 and pages[1]['page']==2
    assert 'Hybrid search handbook' in text
    assert main.chunk_pages(pages,900,120)[1]['page']==2


def test_full_document_is_summarized_once_and_only_summary_is_embedded(client,providers):
    configure(client)
    content=('OriginalOnlyToken '+ 'Full document content. '*2000).encode()
    doc=add_doc(client,content=content)
    main.ingest(doc)
    d=client.get('/api/documents/'+doc).json()
    assert d['status']=='ready'
    assert d['indexing_mode']=='document_summary'
    assert len(d['chunks'])==1 and d['chunks'][0]['page'] is None
    assert d['chunks'][0]['text']==d['summary']
    summaries=[body for url,body in providers if url.endswith('/v1/messages')]
    assert len(summaries)==1
    assert summaries[0]['messages'][0]['content']==content.decode().strip()
    embeddings=[body for url,body in providers if url.endswith('/embed')]
    assert len(embeddings)==1 and embeddings[0]['inputs']==[{'text':d['summary']}]
    upserts=[body for url,body in providers if url.endswith('/vectors/upsert')]
    assert len(upserts)==1 and len(upserts[0]['vectors'])==1
    assert upserts[0]['vectors'][0]['metadata']['text']==d['summary']
    assert not client.post('/api/search',json={'query':'OriginalOnlyToken','mode':'lexical','rerank':False}).json()['results']
